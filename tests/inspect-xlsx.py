"""Independent OOXML verification; the workbook is authored by the TypeScript exporter."""
import sys
import warnings
import zipfile
from datetime import datetime
from math import isclose
from xml.etree import ElementTree
from openpyxl import load_workbook

path = sys.argv[1]
with zipfile.ZipFile(path) as archive:
    assert archive.testzip() is None
    for name in archive.namelist():
        ElementTree.fromstring(archive.read(name))

with warnings.catch_warnings(record=True) as caught:
    warnings.simplefilter("always")
    workbook = load_workbook(path, data_only=False)
assert not caught, [str(w.message) for w in caught]
assert workbook.sheetnames == ["Produção   Sul", "Produção   Sul (2)", "Vazio"]
sheet = workbook.worksheets[0]
assert sheet["A1"].value == "CaféGestor · Safra 2026"
assert sheet["A2"].value == "João & Ana <Sítio> — MG"
assert sheet["A5"].value == "João d'Água & irmãos"
assert sheet["B5"].value == 4990.9 and sheet["B5"].data_type == "n"
assert sheet["B6"].value == -12.5
assert sheet["C6"].value == 0
assert sheet["D5"].value == 37.125
assert isclose(sheet["E5"].value, 0.34125)
assert sheet["F5"].value == datetime(2026, 10, 19)
assert sheet["F6"].value == datetime(2024, 2, 29)
assert sheet["G5"].value is True and sheet["G6"].value is False
assert sheet["B7"].value is None
assert sheet["A6"].value == '=HYPERLINK("https://evil.example","abrir")'
assert sheet["A6"].data_type == "s"
assert sheet["A7"].value == "  @SUM(A1)\nCafé ☕"
assert sheet["B5"].number_format.startswith('"R$" #,##0.00')
assert sheet["C5"].number_format == "#,##0"
assert sheet["D5"].number_format == "#,##0.00"
assert sheet["E5"].number_format == "0.0%"
assert sheet["F5"].number_format == "dd/mm/yyyy"
assert sheet["A4"].fill.fgColor.rgb == "FF163D2B"
assert sheet["A4"].font.color.rgb == "FFFFFFFF"
assert sheet.freeze_panes == "A5"
assert sheet.auto_filter.ref == "A4:G7"
assert sheet.sheet_view.showGridLines is False
assert "A1:G1" in sheet.merged_cells
assert sheet.column_dimensions["A"].width == 40
assert workbook.worksheets[1]["AA2"].value == "AA"
assert workbook.worksheets[1]["AB2"].value == 28
assert workbook.worksheets[2].auto_filter.ref == "A1:A1"
for current in workbook:
    for row in current:
        assert all(cell.data_type != "f" for cell in row), "User text became a formula"
workbook.close()
print("openpyxl: workbook verified; Unicode, types, dates, styles, filters and formula safety passed")
